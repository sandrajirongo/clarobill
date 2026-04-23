"""Report generator — converts analysis results into PDF and DOCX documents.

Produces two outputs:
  1. Full analysis report — branded PDF (findings, line items, AI summary)
  2. Dispute letter — editable DOCX the patient fills in and sends

HTML → PDF via WeasyPrint (synchronous, run in executor by the caller).
Dispute letter → DOCX via python-docx (fast, in-memory).

Prerequisite: A 'reports' Supabase Storage bucket must exist and be private.
Create it in the Supabase dashboard: Storage → New Bucket → name: 'reports', private: true.
"""

import logging
import re
from decimal import Decimal
from io import BytesIO

from app.models.bill import BillingError
from app.services.parser import ParsedLineItem

logger = logging.getLogger(__name__)

_SEVERITY_BADGE = {
    "critical": ("CRITICAL", "#991b1b", "#fee2e2"),
    "warning": ("WARNING", "#92400e", "#fef3c7"),
    "info": ("INFO", "#1e40af", "#dbeafe"),
}

# Matches lines that begin a dispute section in the Claude plain-text output
_DISPUTE_HEADER_RE = re.compile(r"^Dispute\s+\d+\s*:", re.IGNORECASE)


class ReportGenerator:
    """Builds the branded analysis report PDF and the editable dispute letter DOCX."""

    def generate_report_pdf(
        self,
        errors: list[BillingError],
        line_items: list[ParsedLineItem],
        bill_file_name: str,
        ai_summary: str,
    ) -> bytes:
        """Generate the full analysis report as a branded PDF.

        This method is synchronous (WeasyPrint) — the caller must run it in a
        thread pool executor so it does not block the asyncio event loop.
        """
        html = self._build_report_html(errors, line_items, bill_file_name, ai_summary)
        return self._html_to_pdf(html)

    def generate_dispute_letter_docx(self, dispute_letter_text: str) -> bytes:
        """Convert the plain-text dispute letter to an editable Word document (.docx).

        The letter is parsed into paragraphs split on blank lines. Lines matching
        'Dispute N:' are rendered as bold section headers so the patient can navigate
        each dispute item easily before editing.
        """
        from docx import Document  # noqa: PLC0415
        from docx.shared import Pt, RGBColor  # noqa: PLC0415
        from docx.oxml.ns import qn  # noqa: PLC0415
        from docx.oxml import OxmlElement  # noqa: PLC0415

        doc = Document()

        # Page margins
        for section in doc.sections:
            from docx.shared import Inches  # noqa: PLC0415
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.25)

        # Default body style
        style = doc.styles["Normal"]
        style.font.name = "Georgia"
        style.font.size = Pt(11)

        def _add_para(text: str, bold: bool = False, space_before: int = 0, space_after: int = 10) -> None:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(space_before)
            p.paragraph_format.space_after = Pt(space_after)
            p.paragraph_format.line_spacing = Pt(14)
            run = p.add_run(text)
            run.font.name = "Georgia"
            run.font.size = Pt(11)
            run.bold = bold
            if bold:
                run.font.color.rgb = RGBColor(0x2D, 0x4A, 0x2A)  # sage

        # Split on blank lines → paragraphs
        paragraphs = [blk.strip() for blk in dispute_letter_text.split("\n\n") if blk.strip()]

        for block in paragraphs:
            # Each block may contain multiple lines (e.g., recipient address block)
            lines = [ln.strip() for ln in block.splitlines() if ln.strip()]
            if not lines:
                continue

            first_line = lines[0]
            is_dispute_header = bool(_DISPUTE_HEADER_RE.match(first_line))

            if is_dispute_header:
                # "Dispute N:" line bold, rest of block as regular paragraph
                _add_para(first_line, bold=True, space_before=8, space_after=4)
                body = " ".join(lines[1:])
                if body:
                    _add_para(body, space_after=10)
            elif len(lines) > 1:
                # Multi-line block (address block, closing) — one paragraph per line
                for ln in lines:
                    _add_para(ln, space_after=2)
                # Extra gap after the block
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            else:
                _add_para(first_line, space_after=10)

        buf = BytesIO()
        doc.save(buf)
        logger.info("Generated dispute letter DOCX (%d bytes)", buf.tell())
        return buf.getvalue()

    def _build_report_html(
        self,
        errors: list[BillingError],
        line_items: list[ParsedLineItem],
        bill_file_name: str,
        ai_summary: str,
    ) -> str:
        total_charges = sum(Decimal(str(item.total_charge)) for item in line_items)
        total_savings = sum(e.estimated_overcharge for e in errors)

        td = "padding:9px 14px;border-bottom:1px solid #e8e8e8;font-size:11px;color:#1a2e18;vertical-align:top"
        th = "padding:8px 14px;border-bottom:2px solid #2D4A2A;font-size:9px;text-transform:uppercase;letter-spacing:0.6px;color:#2D4A2A;text-align:left;font-weight:600"

        error_rows = "".join(
            f"""<tr>
              <td style="{td};font-weight:600">{e.code.replace('_', ' ').title()}</td>
              <td style="{td}">{e.severity.capitalize()}</td>
              <td style="{td}">{e.description}</td>
              <td style="{td};text-align:right;font-weight:700">${e.estimated_overcharge:.2f}</td>
            </tr>"""
            for e in errors
        )

        line_item_rows = "".join(
            f"""<tr>
              <td style="{td};font-weight:600">{item.cpt_code}</td>
              <td style="{td}">{item.description[:60]}</td>
              <td style="{td}">{item.service_date}</td>
              <td style="{td};text-align:right">${item.total_charge:.2f}</td>
              <td style="{td};color:#6b7c6a">{item.modifier or '—'}</td>
            </tr>"""
            for item in line_items
        )

        findings_section = (
            f"""<table style="width:100%;border-collapse:collapse;margin-top:14px">
              <thead><tr>
                <th style="{th}">Finding</th>
                <th style="{th}">Severity</th>
                <th style="{th}">Description</th>
                <th style="{th};text-align:right">Est. Overcharge</th>
              </tr></thead>
              <tbody>{error_rows}</tbody>
            </table>"""
            if errors else
            '<p style="font-size:12px;color:#2D4A2A;margin-top:14px;font-weight:600">No billing errors were detected.</p>'
        )

        return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <style>
    /* System fonts only — no external @import so WeasyPrint renders without network I/O */
    * {{ box-sizing: border-box; margin: 0; padding: 0; }}
    body {{ font-family: Arial, 'Liberation Sans', sans-serif; color: #1a2e18;
            background: white; font-size: 12px; line-height: 1.65; padding: 56px 64px; }}
    .letterhead {{ display: flex; justify-content: space-between; align-items: flex-end;
                   padding-bottom: 18px; border-bottom: 3px solid #2D4A2A; margin-bottom: 32px; }}
    .logo {{ font-family: Georgia, 'Liberation Serif', serif; font-size: 26px; font-weight: 700;
             font-style: italic; letter-spacing: -0.5px; color: #2D4A2A; }}
    .logo-accent {{ color: #C8F060; }}
    .doc-meta {{ text-align: right; font-size: 10px; color: #6b7c6a; line-height: 1.5; }}
    h2 {{ font-family: Georgia, 'Liberation Serif', serif; font-size: 13px; font-weight: 700;
          color: #2D4A2A; text-transform: uppercase; letter-spacing: 0.08em;
          margin: 36px 0 12px; }}
    .metrics {{ display: flex; gap: 48px; margin-bottom: 8px; }}
    .metric-val {{ font-family: Georgia, 'Liberation Serif', serif; font-size: 22px;
                   font-weight: 700; color: #2D4A2A; line-height: 1.1; }}
    .metric-label {{ font-size: 9px; color: #6b7c6a; text-transform: uppercase;
                     letter-spacing: 0.7px; margin-top: 3px; }}
    .summary-text {{ font-size: 12px; color: #1a2e18; line-height: 1.75; }}
    .footer {{ margin-top: 56px; padding-top: 14px; border-top: 1px solid #ddd;
               font-size: 9px; color: #888; line-height: 1.6; }}
    .footer-brand {{ font-family: Georgia, 'Liberation Serif', serif; font-size: 10px;
                     color: #2D4A2A; font-style: italic; margin-bottom: 5px; }}
  </style>
</head>
<body>

  <div class="letterhead">
    <div class="logo"><span class="logo-accent">Claro</span>Bill</div>
    <div class="doc-meta">
      Medical Bill Analysis Report<br>
      {bill_file_name}
    </div>
  </div>

  <div class="metrics">
    <div>
      <div class="metric-val">${total_charges:.2f}</div>
      <div class="metric-label">Total Charges</div>
    </div>
    <div>
      <div class="metric-val">${total_savings:.2f}</div>
      <div class="metric-label">Potential Savings</div>
    </div>
    <div>
      <div class="metric-val">{len(errors)}</div>
      <div class="metric-label">Issues Found</div>
    </div>
  </div>

  <h2>Summary</h2>
  <p class="summary-text">{ai_summary}</p>

  <h2>Findings</h2>
  {findings_section}

  <h2>Line Items</h2>
  <table style="width:100%;border-collapse:collapse;margin-top:14px">
    <thead><tr>
      <th style="{th}">CPT Code</th>
      <th style="{th}">Description</th>
      <th style="{th}">Service Date</th>
      <th style="{th};text-align:right">Charge</th>
      <th style="{th}">Modifier</th>
    </tr></thead>
    <tbody>{line_item_rows}</tbody>
  </table>

  <div class="footer">
    <div class="footer-brand">ClaroBill</div>
    Prepared using verified Medicare fee schedule data and AI-assisted billing analysis.
    ClaroBill cross-references your charges against published hospital prices, Medicare
    benchmark rates, and NCCI coding rules to identify errors and overcharges.
    Use this report to support a billing dispute with your provider or insurer.
  </div>

</body>
</html>"""

    @staticmethod
    def _html_to_pdf(html: str) -> bytes:
        from weasyprint import HTML  # noqa: PLC0415
        pdf_bytes = HTML(string=html).write_pdf()
        logger.info("Generated report PDF (%d bytes)", len(pdf_bytes))
        return pdf_bytes
